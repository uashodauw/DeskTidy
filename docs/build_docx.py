"""把 docs/ 下的 Markdown 文档转成 Word：python3 docs/build_docx.py 文件名.md"""
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)


def build(md_path: Path):
    docx_path = md_path.with_suffix(".docx")
    subprocess.run(
        ["pandoc", "-f", "markdown-smart", str(md_path), "-o", str(docx_path)],
        check=True,
    )
    doc = Document(docx_path)
    for style in doc.styles:
        if style.type in (1, 2):
            rpr = style.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.insert(0, fonts)
            fonts.set(qn("w:eastAsia"), "PingFang SC")
    for table in doc.tables:
        add_borders(table)
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
    doc.save(docx_path)
    print(f"已生成 {docx_path}")


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        build(Path(arg))
